from flask import Flask, render_template, request, redirect, url_for, jsonify
from apscheduler.schedulers.background import BackgroundScheduler
from flask import send_file
from flask_mysqldb import MySQL
import datetime
import pymysql
import io
import openpyxl
import logging
import re
import os

app = Flask(__name__, template_folder='templates', static_folder='static')


# Configure file logging
logging.basicConfig(
    filename="guesthouse_errors.log",
    level=logging.ERROR,
    format="%(asctime)s [%(levelname)s] %(message)s",
)


def auto_checkout_guests():
    con = None
    try:
        con = connect_db()
        cur = con.cursor()

        now = datetime.datetime.now()

        # 🔍 Get all guest_ids where any row's planned checkout has passed
        cur.execute(
            """
            SELECT DISTINCT guest_id
            FROM bookings
            WHERE
                planned_checkout_date IS NOT NULL
                AND planned_checkout_time IS NOT NULL
                AND checkout_date IS NULL AND checkout_time IS NULL
                AND (
                    planned_checkout_date < %s OR
                    (planned_checkout_date = %s AND planned_checkout_time <= %s)
                )
            """,
            (now.date(), now.date(), now.time()),
        )

        guests_to_checkout = cur.fetchall()

        for (guest_id,) in guests_to_checkout:
            # ✅ Free all beds
            cur.execute(
                "UPDATE beds SET occupied = FALSE, guest_id = NULL WHERE guest_id = %s",
                (guest_id,),
            )

            # ✅ Mark ALL active bookings for that guest
            cur.execute(
                """
                UPDATE bookings
                SET checkout_date = %s, checkout_time = %s
                WHERE guest_id = %s AND checkout_date IS NULL AND checkout_time IS NULL
                """,
                (now.date(), now.time(), guest_id),
            )

        con.commit()

    except Exception as e:
        if con:
            con.rollback()
        print("❌ Error during auto-checkout:", e)

    finally:
        if con:
            con.close()


# Schedule the job every minute
scheduler = BackgroundScheduler()
scheduler.add_job(func=auto_checkout_guests, trigger="interval", minutes=1)
scheduler.start()


app = Flask(__name__)
# Use environment variables for database config (safe for deployment)
app.config["MYSQL_HOST"] = os.environ.get("MYSQL_HOST")
app.config["MYSQL_USER"] = os.environ.get("MYSQL_USER")
app.config["MYSQL_PASSWORD"] = os.environ.get("MYSQL_PASSWORD")
app.config["MYSQL_DB"] = os.environ.get("MYSQL_DB")

mysql = MySQL(app)
app.secret_key = "your_secret_key"
meal_prices = {"Breakfast": 50, "Lunch": 100, "Dinner": 80}


def connect_db():
    con = pymysql.connect(
        host=os.environ.get("MYSQL_HOST"),
        user=os.environ.get("MYSQL_USER"),
        password=os.environ.get("MYSQL_PASSWORD"),
        database=os.environ.get("MYSQL_DB"),
        port=int(os.environ.get("MYSQL_PORT", 3306)),
        cursorclass=pymysql.cursors.Cursor,
    )
    cur = con.cursor()
    cur.execute("SET SESSION innodb_lock_wait_timeout = 10")
    return con
    

@app.route("/")
# def index():
#     return "✅ Flask app is running!"
def index():
    print("➡️ Reached index route")
    con = connect_db()
    cur = con.cursor()

    total_beds = 6 * 2  # 6 rooms, 2 beds each = 12 beds total

    # Count how many beds are occupied
    cur.execute("SELECT COUNT(*) FROM beds WHERE occupied=TRUE")
    occupied = cur.fetchone()[0]
    available = total_beds - occupied
    # occupied = cur.fetchone()[0]
    # available = total_beds - occupied

    # Get all booking logs (past + current bookings)
    cur.execute(
        """
        SELECT g.id, g.name, b.room_number, b.bed_number, 
               bk.checkin_date, bk.checkin_time, 
               bk.checkout_date, bk.checkout_time, 
               g.reason, g.category, g.meals, g.meal_status
        FROM bookings bk
        JOIN guests g ON bk.guest_id = g.id
        JOIN beds b ON bk.bed_id = b.id
        ORDER BY bk.checkin_date DESC, bk.checkin_time DESC
        """
    )
    logs = cur.fetchall()

    # Fetch only the latest check-in per currently occupied guest
    cur.execute(
        """
        SELECT g.id, g.name, b.room_number, b.bed_number, bk.checkin_date, bk.checkin_time, g.category
        FROM bookings bk
        JOIN guests g ON bk.guest_id = g.id
        JOIN beds b ON bk.bed_id = b.id
        WHERE b.occupied = TRUE
        AND bk.checkout_date IS NULL  -- Ensures we only include guests who have not checked out
        AND bk.booking_id = (
            SELECT MAX(bk2.booking_id)
            FROM bookings bk2
            WHERE bk2.guest_id = bk.guest_id
        )
        ORDER BY bk.checkin_date ASC, bk.checkin_time ASC;
        """
    )
    current_guests = cur.fetchall()  # Only the latest booking per occupied guest

    # Count how many beds each guest booked
    cur.execute(
        """
        SELECT g.id, COUNT(*) as bed_count
        FROM bookings bk
        JOIN guests g ON bk.guest_id = g.id
        WHERE bk.checkout_date IS NULL
        GROUP BY g.id
    """
    )
    bed_counts_per_guest = dict(cur.fetchall())

    # cur.execute(
    #     """
    # SELECT g.id AS guest_id, m.name, m.message, m.room_number, m.bed_number, m.timestamp
    # FROM guest_messages m
    # JOIN guests g ON g.name = m.name
    # JOIN bookings bk ON g.id = bk.guest_id
    # JOIN beds b ON bk.bed_id = b.id
    # WHERE b.room_number = m.room_number
    #   AND b.bed_number = m.bed_number
    #   AND bk.checkout_date IS NULL
    # ORDER BY m.timestamp DESC
    # """
    # )
    # messages = cur.fetchall()

    cur.execute(
        """
    SELECT g.id AS guest_id, m.id AS message_id, m.name, m.message, m.room_number, m.bed_number, m.timestamp
    FROM guest_messages m
    JOIN guests g ON g.name = m.name
    JOIN bookings bk ON g.id = bk.guest_id
    JOIN beds b ON bk.bed_id = b.id
    WHERE b.room_number = m.room_number
      AND b.bed_number = m.bed_number
      AND bk.checkout_date IS NULL
    ORDER BY m.timestamp DESC
    """
    )

    messages = cur.fetchall()

    cur.execute(
        "SELECT id, title, message, posted_on FROM announcements ORDER BY posted_on DESC"
    )
    announcements = cur.fetchall()

    # # Fetch unique meal orders (one per guest)
    #     cur.execute("SELECT id, name, meals FROM guests WHERE meals IS NOT NULL AND meals != ''")
    #     meal_orders = cur.fetchall()

    con.close()
    try:
        return render_template(
            "index.html",
            logs=logs,
            current_guests=current_guests,
            total_beds=total_beds,
            occupied=occupied,
            available=available,
            bed_counts=bed_counts_per_guest,
            messages=messages,
            announcements=announcements,
            meal_prices=meal_prices,
        )
    except Exception as e:
        print("🔥 Template render failed:", str(e))
        return "Template error: " + str(e)



@app.route("/register", methods=["GET", "POST"])
def register_guest():
    con = connect_db()
    cur = con.cursor()

    if request.method == "POST":
        name = request.form["name"]
        email = request.form["email"]
        phone = request.form["phone"]
        reason = request.form["reason"]
        category = request.form["category"]
        meals = request.form.getlist("meals")  # ['Breakfast', 'Lunch', 'Dinner']
        meal_summary = ", ".join(meals)  # 'Breakfast, Lunch, Dinner'

        # ✅ Get selected beds from form
        selected_bed_ids = request.form.getlist("selected_beds")

        if not selected_bed_ids:
            con.close()
            return "Error: You must select at least one available bed."

        # ✅ Verify all selected beds are still available
        placeholders = ",".join(["%s"] * len(selected_bed_ids))
        cur.execute(
            f"SELECT id FROM beds WHERE id IN ({placeholders}) AND occupied = FALSE",
            selected_bed_ids,
        )
        available = cur.fetchall()

        if len(available) != len(selected_bed_ids):
            con.close()
            return "Error: One or more selected beds are no longer available."

        # ✅ Insert guest
        cur.execute(
            "INSERT INTO guests (name, email, phone, reason, category, meals) VALUES (%s, %s, %s, %s, %s, %s)",
            (name, email, phone, reason, category, meal_summary),
        )
        guest_id = cur.lastrowid

        checkin_date = datetime.datetime.now().date()
        checkin_time = datetime.datetime.now().time()

        # ✅ Assign each selected bed
        for bed_id in selected_bed_ids:
            cur.execute(
                "INSERT INTO bookings (guest_id, bed_id, checkin_date, checkin_time) VALUES (%s, %s, %s, %s)",
                (guest_id, bed_id, checkin_date, checkin_time),
            )
            cur.execute(
                "UPDATE beds SET occupied = TRUE, guest_id = %s WHERE id = %s",
                (guest_id, bed_id),
            )

        con.commit()
        con.close()
        return redirect("/")

    # GET method – fetch available beds
    cur.execute(
        "SELECT id, room_number, bed_number, occupied FROM beds ORDER BY room_number, bed_number"
    )

    all_beds = cur.fetchall()

    cur.execute("SELECT COUNT(*) FROM beds WHERE occupied = FALSE")
    max_available_beds = cur.fetchone()[0]
    con.close()

    return render_template(
        "register_guest.html",
        max_beds=max_available_beds,
        available_beds=max_available_beds,
        all_beds=all_beds,
    )


@app.route("/remove_guest/<int:guest_id>", methods=["POST"])
def remove_guest(guest_id):
    con = connect_db()
    cur = con.cursor()

    # Step 1: Free the bed
    cur.execute(
        "UPDATE beds SET occupied = FALSE, guest_id = NULL WHERE guest_id = %s",
        (guest_id,),
    )

    # Log check-out in bookings table
    checkout_date = datetime.datetime.now().date()
    checkout_time = datetime.datetime.now().time()

    cur.execute(
        """
        UPDATE bookings
        SET checkout_date = %s, checkout_time = %s
        WHERE guest_id = %s AND checkout_date IS NULL
    """,
        (checkout_date, checkout_time, guest_id),
    )

    # Step 2: Delete the guest
    # cur.execute("DELETE FROM guests WHERE id = %s", (guest_id,))

    con.commit()
    con.close()

    return redirect("/")


# ONLY WORKING ON ONE ROW AUTO CHECK OUT
# @app.route("/submit", methods=["GET", "POST"])
# def submit_form():
#     con = None
#     try:
#         if request.method == "POST":
#             name = request.form["name"]
#             email = request.form["email"]
#             phone = request.form["phone"]
#             reason = request.form["reason"]
#             category = request.form["category"]
#             meals = request.form.getlist("meals")
#             meal_summary = ", ".join(meals)
#             adults = request.form["adults"]
#             childs = request.form["childs"]
#             planned_checkout_date = request.form["checkout_date"]
#             planned_checkout_time = request.form["checkout_time"]

#             selected_bed_ids = request.form.getlist("selected_beds")

#             if not selected_bed_ids:
#                 return "Error: You must select at least one available bed."

#             con = connect_db()
#             cur = con.cursor()

#             placeholders = ",".join(["%s"] * len(selected_bed_ids))
#             cur.execute(
#                 f"SELECT id FROM beds WHERE id IN ({placeholders}) AND occupied = FALSE",
#                 selected_bed_ids,
#             )
#             available = cur.fetchall()

#             if len(available) != len(selected_bed_ids):
#                 return "Error: One or more selected beds are no longer available."

#             cur.execute(
#                 "INSERT INTO guests (name, email, phone, reason, category, meals, adults, childs) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)",
#                 (name, email, phone, reason, category, meal_summary, adults, childs),
#             )
#             guest_id = cur.lastrowid

#             checkin_date = datetime.datetime.now().date()
#             checkin_time = datetime.datetime.now().time()

#             for bed_id in selected_bed_ids:
#                 bed_id = int(bed_id)
#                 cur.execute(
#                     "UPDATE beds SET occupied = TRUE, guest_id = %s WHERE id = %s",
#                     (guest_id, bed_id),
#                 )
#                 cur.execute(
#                     """
#                     INSERT INTO bookings (
#                         guest_id, bed_id, checkin_date, checkin_time,
#                         planned_checkout_date, planned_checkout_time
#                     ) VALUES (%s, %s, %s, %s, %s, %s)
#                     """,
#                     (
#                         guest_id,
#                         bed_id,
#                         checkin_date,
#                         checkin_time,
#                         planned_checkout_date,
#                         planned_checkout_time,
#                     ),
#                 )

#             cur.execute(
#                 "SELECT room_number, bed_number FROM beds WHERE guest_id = %s",
#                 (guest_id,),
#             )
#             assigned_beds = cur.fetchall()

#             con.commit()
#             return render_template(
#                 "welcome.html",
#                 name=name,
#                 guest_id=guest_id,
#                 assigned_beds=assigned_beds,
#             )

#         # ---------- GET method ----------
#         con = connect_db()
#         cur = con.cursor()

#         cur.execute(
#             "SELECT id, room_number, bed_number, occupied FROM beds ORDER BY room_number, bed_number"
#         )
#         all_beds = cur.fetchall()

#         cur.execute("SELECT COUNT(*) FROM beds WHERE occupied = FALSE")
#         max_available_beds = cur.fetchone()[0]
#         no_beds = max_available_beds == 0

#         cur.execute("SELECT DISTINCT room_number FROM beds WHERE occupied = FALSE")
#         available_rooms = [row[0] for row in cur.fetchall()]
#         room_count = len(available_rooms)

#         return render_template(
#             "form.html",
#             max_beds=max_available_beds,
#             available_rooms=available_rooms,
#             room_count=room_count,
#             no_beds=no_beds,
#             all_beds=all_beds,
#         )

#     except Exception as e:
#         if con:
#             con.rollback()
#         print("❌ Error in submit_form:", e)
#         return "An internal error occurred. Please try again."

#     finally:
#         if con:
#             con.close()


@app.route("/submit", methods=["GET", "POST"])
def submit_form():
    if request.method == "POST":
        name = request.form["name"].strip()
        email = request.form["email"].strip()
        phone = request.form["phone"].strip()
        reason = request.form["reason"].strip()
        category = request.form["category"]
        meals = request.form.getlist("meals")
        meal_summary = ", ".join(meals)
        adults = int(request.form["adults"])
        childs = int(request.form["childs"])
        planned_checkout_date = request.form["checkout_date"]
        planned_checkout_time = request.form["checkout_time"]
        selected_bed_ids = list(map(int, request.form.getlist("selected_beds")))

        # ✅ Validation
        if not name or not phone:
            return "❌ Name and phone are required."

        if not re.match(r"[^@]+@[^@]+\.[^@]+", email):
            return "❌ Invalid email address."

        if not any(c.isalnum() for c in reason):
            return "❌ Reason must contain words or valid explanation."

        if not category:
            return "❌ Visiting category is required."

        if adults <= 0 and childs <= 0:
            return "❌ At least one person must be present."

        if not planned_checkout_date or not planned_checkout_time:
            return "❌ Planned checkout date and time are required."

        if not selected_bed_ids:
            return "❌ You must select at least one available bed."

        # ✅ DB logic in try/finally
        con = None
        try:
            con = connect_db()
            cur = con.cursor()

            placeholders = ",".join(["%s"] * len(selected_bed_ids))
            cur.execute(
                f"SELECT id FROM beds WHERE id IN ({placeholders}) AND occupied = FALSE",
                selected_bed_ids,
            )
            available = cur.fetchall()

            if len(available) != len(selected_bed_ids):
                return "❌ One or more selected beds are no longer available."

            # ✅ Insert guest
            cur.execute(
                """
                INSERT INTO guests (name, email, phone, reason, category, meals, adults, childs)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                """,
                (name, email, phone, reason, category, meal_summary, adults, childs),
            )
            guest_id = cur.lastrowid

            checkin_date = datetime.datetime.now().date()
            checkin_time = datetime.datetime.now().time()

            for bed_id in selected_bed_ids:
                cur.execute(
                    "UPDATE beds SET occupied = TRUE, guest_id = %s WHERE id = %s",
                    (guest_id, bed_id),
                )
                cur.execute(
                    """
                    INSERT INTO bookings (
                        guest_id, bed_id, checkin_date, checkin_time,
                        planned_checkout_date, planned_checkout_time,
                        checkout_date, checkout_time
                    ) VALUES (%s, %s, %s, %s, %s, %s, NULL, NULL)
                    """,
                    (
                        guest_id,
                        bed_id,
                        checkin_date,
                        checkin_time,
                        planned_checkout_date,
                        planned_checkout_time,
                    ),
                )

            # ✅ Fetch bed info
            cur.execute(
                "SELECT room_number, bed_number FROM beds WHERE guest_id = %s",
                (guest_id,),
            )
            assigned_beds = cur.fetchall()

            con.commit()

        except Exception as e:
            if con:
                con.rollback()
            return f"❌ Error: {str(e)}"

        finally:
            if con:
                con.close()

        return render_template(
            "welcome.html",
            name=name,
            guest_id=guest_id,
            assigned_beds=assigned_beds,
        )

    # ---------- GET: Render form ----------
    con = connect_db()
    cur = con.cursor()

    cur.execute(
        "SELECT id, room_number, bed_number, occupied FROM beds ORDER BY room_number, bed_number"
    )
    all_beds = cur.fetchall()

    cur.execute("SELECT COUNT(*) FROM beds WHERE occupied = FALSE")
    max_available_beds = cur.fetchone()[0]
    no_beds = max_available_beds == 0

    cur.execute("SELECT DISTINCT room_number FROM beds WHERE occupied = FALSE")
    available_rooms = [row[0] for row in cur.fetchall()]
    room_count = len(available_rooms)

    con.close()

    return render_template(
        "form.html",
        max_beds=max_available_beds,
        available_rooms=available_rooms,
        room_count=room_count,
        no_beds=no_beds,
        all_beds=all_beds,
    )


@app.route("/welcome")
def welcome():
    name = request.args.get("name", "Guest")
    return render_template("welcome.html", name=name)


@app.route("/mark_meal_complete/<int:guest_id>", methods=["POST"])
def mark_meal_complete(guest_id):
    con = connect_db()
    cur = con.cursor()
    cur.execute(
        "UPDATE guests SET meal_status = 'Completed' WHERE id = %s", (guest_id,)
    )
    con.commit()
    con.close()
    return redirect("/#guests")  # or use url_for("index") if preferred


@app.route("/guest")
def guest_ui():
    con = connect_db()
    cur = con.cursor()

    # Total beds available
    cur.execute("SELECT COUNT(*) FROM beds WHERE occupied = FALSE")
    max_available_beds = cur.fetchone()[0]
    no_beds = max_available_beds == 0

    # Available room numbers
    cur.execute("SELECT DISTINCT room_number FROM beds WHERE occupied = FALSE")
    rooms = cur.fetchall()
    available_rooms = [row[0] for row in rooms]
    room_count = len(available_rooms)

    # Fetch announcements
    cur.execute(
        "SELECT title, message, posted_on FROM announcements ORDER BY posted_on DESC"
    )
    announcements = cur.fetchall()

    # ✅ Fetch all beds for visual selection
    cur.execute(
        "SELECT id, room_number, bed_number, occupied FROM beds ORDER BY room_number, bed_number"
    )
    all_beds = cur.fetchall()

    con.close()

    return render_template(
        "form.html",
        max_beds=max_available_beds,
        available_beds=max_available_beds,
        available_rooms=available_rooms,
        room_count=room_count,
        no_beds=no_beds,
        announcements=announcements,
        all_beds=all_beds,  # ✅ pass to form.html
    )


@app.route("/message", methods=["POST"])
def guest_message():
    guest_id = int(request.form["guest_id"])
    room_number = int(request.form["room_number"])
    bed_number = int(request.form["bed_number"])
    message = request.form["message"]

    con = connect_db()
    cur = con.cursor()

    # ✅ Strict validation: guest must be in an occupied bed that matches
    cur.execute(
        """
        SELECT g.name
        FROM guests g
        JOIN bookings bk ON g.id = bk.guest_id
        JOIN beds b ON bk.bed_id = b.id
        WHERE g.id = %s AND b.room_number = %s AND b.bed_number = %s
          AND b.occupied = TRUE AND bk.checkout_date IS NULL
        ORDER BY bk.booking_id DESC
        LIMIT 1
    """,
        (guest_id, room_number, bed_number),
    )

    result = cur.fetchone()
    if not result:
        con.close()
        return """
        <div style='
          background-color: #ffdddd;
          border-left: 6px solid #f44336;
          padding: 16px;
          margin: 40px auto;
          max-width: 600px;
          font-family: Arial, sans-serif;
        '>
          <h3>⚠️ Error: Invalid Details</h3>
          <p>The Guest ID, Room Number, and Bed Number combination does not match any current booking.</p>
          <a href='/guest'>← Go Back</a>
        </div>
        """

    name = result[0]

    cur.execute(
        """
        INSERT INTO guest_messages (name, message, room_number, bed_number)
        VALUES (%s, %s, %s, %s)
    """,
        (name, message, room_number, bed_number),
    )

    con.commit()
    con.close()
    return redirect(url_for("req_success"))


@app.route("/req_success")
def req_success():
    return render_template("req_success.html")


@app.route("/delete_message/<int:message_id>", methods=["POST"])
def delete_message(message_id):
    con = connect_db()
    cur = con.cursor()
    cur.execute("DELETE FROM guest_messages WHERE id = %s", (message_id,))
    con.commit()
    con.close()
    return redirect("/#maintenance")  # Redirect back to the messages section


@app.route("/guest_rooms_beds/<int:guest_id>")
def guest_rooms_beds(guest_id):
    con = connect_db()
    cur = con.cursor()
    cur.execute(
        """
        SELECT b.room_number, b.bed_number
        FROM bookings bk
        JOIN beds b ON bk.bed_id = b.id
        WHERE bk.guest_id = %s AND bk.checkout_date IS NULL AND b.occupied = TRUE;
    """,
        (guest_id,),
    )
    data = [{"room_number": r, "bed_number": b} for r, b in cur.fetchall()]
    con.close()
    print(data)  # Add this line for debug
    return jsonify(data)


@app.route("/add_announcement", methods=["POST"])
def add_announcement():
    title = request.form["title"]
    message = request.form["message"]

    con = connect_db()
    cur = con.cursor()
    cur.execute(
        "INSERT INTO announcements (title, message) VALUES (%s, %s)", (title, message)
    )
    con.commit()
    con.close()
    return redirect("/#announcements")


@app.route("/delete_announcement/<int:announcement_id>", methods=["POST"])
def delete_announcement(announcement_id):
    con = connect_db()
    cur = con.cursor()
    cur.execute("DELETE FROM announcements WHERE id = %s", (announcement_id,))
    con.commit()
    con.close()
    return redirect("/#announcements")


@app.route("/edit_announcement/<int:announcement_id>", methods=["GET", "POST"])
def edit_announcement(announcement_id):
    con = connect_db()
    cur = con.cursor()

    if request.method == "POST":
        title = request.form["title"]
        message = request.form["message"]
        cur.execute(
            "UPDATE announcements SET title = %s, message = %s WHERE id = %s",
            (title, message, announcement_id),
        )
        con.commit()
        con.close()
        return redirect("/#announcements")

    # GET request – show edit form
    cur.execute(
        "SELECT title, message FROM announcements WHERE id = %s", (announcement_id,)
    )
    announcement = cur.fetchone()
    con.close()
    return render_template(
        "edit_announcement.html",
        id=announcement_id,
        title=announcement[0],
        message=announcement[1],
    )


@app.route("/export_logs")
def export_logs():
    con = connect_db()
    cur = con.cursor()

    # Get all booking logs (past + current)
    cur.execute(
        """
        SELECT g.id, g.name, b.room_number, b.bed_number, 
               bk.checkin_date, bk.checkin_time, 
               bk.checkout_date, bk.checkout_time, 
               g.reason, g.category, g.meals, g.adults, g.childs,
               (g.adults + g.childs) AS total_persons
        FROM bookings bk
        JOIN guests g ON bk.guest_id = g.id
        JOIN beds b ON bk.bed_id = b.id
        ORDER BY bk.checkin_date DESC, bk.checkin_time DESC
        """
    )
    logs = cur.fetchall()
    con.close()

    # Create Excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Guest Logs"

    # Write headers
    headers = [
        "Guest ID",
        "Name",
        "Room",
        "Bed",
        "Check-In Date",
        "Check-In Time",
        "Check-Out Date",
        "Check-Out Time",
        "Purpose",
        "Visiting Plant",
        "Meals",
        "Adults",
        "Children",
        "Total Persons",
    ]
    ws.append(headers)

    # Write data rows
    for row in logs:
        ws.append(list(row))

    # Save to in-memory file
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        download_name="guest_logs.xlsx",
        as_attachment=True,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/clear_logs", methods=["POST"])
def clear_logs():
    try:
        con = connect_db()
        cur = con.cursor()

        # Delete only bookings where guest has checked out
        cur.execute("DELETE FROM bookings WHERE checkout_date IS NOT NULL")

        con.commit()
        return redirect(url_for("index"))

    except Exception as e:
        if con:
            con.rollback()
        return f"❌ Error while clearing logs: {str(e)}"

    finally:
        if con:
            con.close()


@app.route("/check_updates")
def check_updates():
    connection = connect_db()  # Use the function
    cursor = connection.cursor()

    cursor.execute("SELECT COUNT(*) FROM guests")
    record_count = cursor.fetchone()[0]

    cursor.close()
    connection.close()

    return jsonify({"total_records": record_count})


@app.route("/expense")
def expense():
    return render_template("chart.html")


@app.route('/')
def home():
    return "✅ Flask is working!"

@app.route('/test-db')
def test_db():
    try:
        with conn.cursor() as cursor:
            cursor.execute("SELECT NOW()")
            now = cursor.fetchone()
        return f"MySQL connected. Time: {now[0]}"
    except Exception as e:
        return f"❌ DB Error: {str(e)}"


pass
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
    
